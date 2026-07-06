# q_excel/disk.py

from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from q_excel.utils import (
    clean_dataframe,
    dataframe_to_dicts,
    get_file_suffix,
)


def read_file_from_disk(
    file_path: Union[str, Path],
    has_header: bool = True,
    sheet_name: Union[int, str] = 0,
    csv_encoding: str = "utf-8",
) -> pd.DataFrame:
    """
    Læser en fil fra disk.

    file_path:
        Konkret sti til filen.

    Denne funktion opretter ingen mapper.
    Processen skal selv give den rigtige sti.
    """

    file_path_object = Path(file_path)
    file_suffix = get_file_suffix(str(file_path_object))

    if file_suffix == ".xlsx":
        return read_excel_from_disk(
            file_path=file_path_object,
            has_header=has_header,
            sheet_name=sheet_name,
        )

    if file_suffix == ".csv":
        return read_csv_from_disk(
            file_path=file_path_object,
            has_header=has_header,
            encoding=csv_encoding,
        )

    raise ValueError(f"Unsupported file type: {file_suffix}")


def read_file_from_disk_as_dicts(
    file_path: Union[str, Path],
    has_header: bool = True,
    sheet_name: Union[int, str] = 0,
    csv_encoding: str = "utf-8",
) -> List[Dict[str, Any]]:
    """
    Læser en fil fra disk og returnerer data som dicts.

    Denne funktion opretter ingen mapper.
    """

    df = read_file_from_disk(
        file_path=file_path,
        has_header=has_header,
        sheet_name=sheet_name,
        csv_encoding=csv_encoding,
    )

    return dataframe_to_dicts(df)


def read_excel_from_disk(
    file_path: Union[str, Path],
    has_header: bool = True,
    sheet_name: Union[int, str] = 0,
) -> pd.DataFrame:
    """
    Læser Excel-fil fra disk.

    file_path skal være en konkret sti.
    """

    file_path_object = Path(file_path)

    if has_header:
        df = pd.read_excel(
            file_path_object,
            dtype=str,
            engine="openpyxl",
            sheet_name=sheet_name,
        )
    else:
        df = pd.read_excel(
            file_path_object,
            header=None,
            dtype=str,
            engine="openpyxl",
            sheet_name=sheet_name,
        )
        df.columns = [f"Column_{i + 1}" for i in range(len(df.columns))]

    return clean_dataframe(df)


def read_csv_from_disk(
    file_path: Union[str, Path],
    has_header: bool = True,
    encoding: str = "utf-8",
) -> pd.DataFrame:
    """
    Læser CSV-fil fra disk.

    file_path skal være en konkret sti.
    """

    file_path_object = Path(file_path)

    if has_header:
        df = pd.read_csv(
            file_path_object,
            dtype=str,
            encoding=encoding,
        )
    else:
        df = pd.read_csv(
            file_path_object,
            header=None,
            dtype=str,
            encoding=encoding,
        )
        df.columns = [f"Column_{i + 1}" for i in range(len(df.columns))]

    return clean_dataframe(df)


def excel_cell_edit(
    file_path: Union[str, Path],
    mode: str,
    from_ref: Optional[str] = None,
    to_ref: Optional[str] = None,
    value=None,
    sheet_name: Optional[str] = None,
    save: bool = False,
    save_path: Optional[Union[str, Path]] = None,
):
    """
    Læser, kopierer, skriver eller sletter celler i en Excel-fil.

    Denne funktion arbejder kun med konkrete filstier.

    Den opretter ingen mapper.

    Modes:
        read
        copy
        write
        clear
    """

    file_path_object = Path(file_path)

    wb = load_workbook(file_path_object, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    print(f"Bruger ark: '{ws.title}'")

    # ---------- READ ----------
    if mode == "read":
        if not from_ref:
            raise ValueError("read kræver from_ref")

        if ":" in from_ref:
            min_col, min_row, max_col, max_row = range_boundaries(from_ref)

            data = []
            for row in ws.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
            ):
                data.append([cell.value for cell in row])

            return data

        return ws[from_ref].value

    # ---------- COPY ----------
    if mode == "copy":
        if not from_ref or not to_ref:
            raise ValueError("copy kræver både from_ref og to_ref")

        if ":" in from_ref and ":" in to_ref:
            min_c1, min_r1, max_c1, max_r1 = range_boundaries(from_ref)
            min_c2, min_r2, max_c2, max_r2 = range_boundaries(to_ref)

            rows1 = max_r1 - min_r1 + 1
            cols1 = max_c1 - min_c1 + 1
            rows2 = max_r2 - min_r2 + 1
            cols2 = max_c2 - min_c2 + 1

            if rows1 != rows2 or cols1 != cols2:
                raise ValueError("Ranges skal have samme størrelse")

            for i, row in enumerate(
                ws.iter_rows(
                    min_row=min_r1,
                    max_row=max_r1,
                    min_col=min_c1,
                    max_col=max_c1,
                )
            ):
                for j, cell in enumerate(row):
                    ws.cell(
                        row=min_r2 + i,
                        column=min_c2 + j,
                    ).value = cell.value
        else:
            ws[to_ref] = ws[from_ref].value

    # ---------- WRITE ----------
    elif mode == "write":
        if not to_ref:
            raise ValueError("write kræver to_ref")

        if ":" in to_ref:
            min_col, min_row, max_col, max_row = range_boundaries(to_ref)

            for row in ws.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
            ):
                for cell in row:
                    cell.value = value
        else:
            ws[to_ref] = value

    # ---------- CLEAR ----------
    elif mode == "clear":
        if not to_ref:
            raise ValueError("clear kræver to_ref")

        if ":" in to_ref:
            min_col, min_row, max_col, max_row = range_boundaries(to_ref)

            for row in ws.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
            ):
                for cell in row:
                    cell.value = None
        else:
            ws[to_ref] = None

    else:
        raise ValueError("Invalid mode. Brug read, copy, write eller clear.")

    # ---------- SAVE ----------
    if save:
        if not save_path:
            raise ValueError("save_path skal angives når save=True")

        save_path_object = Path(save_path)

        if not save_path_object.parent.exists():
            raise FileNotFoundError(
                f"Mappen findes ikke: {save_path_object.parent}. "
                "Processen skal selv oprette mapper."
            )

        print(f"Gemmer fil til: {save_path_object}")

        wb.save(save_path_object)

    return None